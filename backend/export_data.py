"""
export_data.py
---------------
Xuất lại dữ liệu điều luật ĐANG CÓ trong Postgres ra file .docx, dùng làm
file "dữ liệu thô" (data/btl_nlp.docx) để chỉnh sửa/thêm mới về sau.

Hữu ích khi bạn lỡ mất file btl_nlp.docx gốc nhưng dữ liệu vẫn còn nguyên
trong database (import_data.py đã từng chạy thành công) - export lại để
có 1 file docx CHÍNH XÁC khớp với dữ liệu đang chạy, làm điểm khởi đầu an
toàn để thêm điều luật mới (qua prepare_data.py + import_data.py sau này).

Chạy (trong container):
    docker compose exec backend python export_data.py

Mặc định ghi ra /app/btl_nlp_export.docx (bên trong container, vì thư mục
/app/data được mount READ-ONLY nên không ghi trực tiếp vào đó được). Sau
khi chạy xong, copy file đó ra máy host bằng lệnh (chạy ở PowerShell,
KHÔNG phải trong container):
    docker cp chatbot_backend:/app/btl_nlp_export.docx ./data/btl_nlp.docx
"""

import argparse
import os

import psycopg2
from docx import Document

DATABASE_URL = os.environ.get(
    "DATABASE_URL", "postgresql://chatbot:chatbot@localhost:5432/luat_db"
)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out",
        default=os.path.join(os.path.dirname(__file__), "btl_nlp_export.docx"),
        help="Đường dẫn file docx sẽ tạo ra (bên trong container).",
    )
    args = parser.parse_args()

    conn = psycopg2.connect(DATABASE_URL)
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT category, content
                FROM legal_documents
                ORDER BY category NULLS FIRST, id
                """
            )
            rows = cur.fetchall()
    finally:
        conn.close()

    print(f"Đã lấy {len(rows)} bản ghi từ database.")

    doc = Document()
    last_category = object()  # chắc chắn khác bất kỳ category thật nào ban đầu
    for category, content in rows:
        if category != last_category:
            if category:
                doc.add_paragraph(category)
            last_category = category
        doc.add_paragraph(content)

    doc.save(args.out)
    print(f"Đã lưu {args.out} ({len(rows)} dòng nội dung).")
    print("Nhớ chạy 'docker cp' để đưa file này ra ngoài host (xem docstring đầu file).")


if __name__ == "__main__":
    main()
